"""
services/data_collection.py
Sammelt Stellenanzeigen aus verschiedenen Quellen

Basiert auf:
- Groß (2022): Text Mining im Personalmanagement
- Wu (2024): Data Mining with Python
"""

import logging
import re
import io
from pathlib import Path
from typing import List, Optional
from datetime import datetime
import hashlib

# PDF/Word Reader
try:
    import PyPDF2
    import pdfplumber
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False
    logging.warning("PDF-Bibliotheken nicht verfügbar")

try:
    from docx import Document
    WORD_AVAILABLE = True
except ImportError:
    WORD_AVAILABLE = False
    logging.warning("Word-Bibliothek nicht verfügbar")

# Google Drive
try:
    from google.oauth2.credentials import Credentials
    from googleapiclient.discovery import build
    from google.auth.transport.requests import Request
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.http import MediaIoBaseDownload
    GDRIVE_AVAILABLE = True
except ImportError:
    GDRIVE_AVAILABLE = False
    logging.warning("Google Drive nicht verfügbar")

from models.job_ad import JobAd, Organization


class FileReader:
    """Liest verschiedene Dateiformate"""

    def __init__(self):
        self.logger = logging.getLogger(__name__)

    def read_pdf(self, file_source, is_bytes=False) -> str:
        """Liest PDF mit Fallback-Strategie"""
        if not PDF_AVAILABLE:
            self.logger.error("PDF-Bibliotheken nicht installiert")
            return ""

        text = ""

        # Versuch 1: pdfplumber (bessere Textextraktion)
        try:
            with pdfplumber.open(file_source) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"

            if text.strip():
                return text
        except Exception as e:
            self.logger.debug(f"pdfplumber fehlgeschlagen: {e}")

        # Versuch 2: PyPDF2 (Fallback)
        try:
            if is_bytes:
                pdf_reader = PyPDF2.PdfReader(file_source)
            else:
                with open(file_source, 'rb') as f:
                    pdf_reader = PyPDF2.PdfReader(f)

            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"

            return text
        except Exception as e:
            self.logger.error(f"PDF-Extraktion fehlgeschlagen: {e}")
            return ""

    def read_word(self, file_source, is_bytes=False) -> str:
        """Liest Word-Dokumente"""
        if not WORD_AVAILABLE:
            self.logger.error("Word-Bibliothek nicht installiert")
            return ""

        try:
            doc = Document(file_source)

            # Paragraphen
            paragraphs = []
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)

            # Tabellen
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            paragraphs.append(cell.text)

            return "\n".join(paragraphs)

        except Exception as e:
            self.logger.error(f"Word-Extraktion fehlgeschlagen: {e}")
            return ""

    def read_file(self, file_path: Path) -> Optional[str]:
        """Liest Datei basierend auf Extension"""
        suffix = file_path.suffix.lower()

        if suffix == '.pdf':
            return self.read_pdf(file_path)
        elif suffix in ['.docx', '.doc']:
            return self.read_word(file_path)
        else:
            self.logger.warning(f"Nicht unterstütztes Format: {suffix}")
            return None


class GoogleDriveConnector:
    """Verbindung zu Google Drive"""

    SCOPES = ['https://www.googleapis.com/auth/drive.readonly']

    def __init__(self, credentials_file='credentials.json', token_file='token.json'):
        self.logger = logging.getLogger(__name__)
        self.credentials_file = Path(credentials_file)
        self.token_file = Path(token_file)
        self.service = None

    def authenticate(self) -> bool:
        """Authentifiziert mit Google Drive"""
        if not GDRIVE_AVAILABLE:
            self.logger.error("Google Drive Bibliotheken nicht installiert")
            return False

        try:
            creds = None

            # Token laden
            if self.token_file.exists():
                creds = Credentials.from_authorized_user_file(
                    str(self.token_file), self.SCOPES
                )

            # Login falls nötig
            if not creds or not creds.valid:
                if creds and creds.expired and creds.refresh_token:
                    creds.refresh(Request())
                else:
                    if not self.credentials_file.exists():
                        self.logger.error(f"credentials.json nicht gefunden")
                        return False

                    flow = InstalledAppFlow.from_client_secrets_file(
                        str(self.credentials_file), self.SCOPES
                    )
                    creds = flow.run_local_server(port=0)

                # Token speichern
                self.token_file.write_text(creds.to_json())

            self.service = build('drive', 'v3', credentials=creds)
            self.logger.info("✅ Google Drive authentifiziert")
            return True

        except Exception as e:
            self.logger.error(f"Google Drive Authentifizierung fehlgeschlagen: {e}")
            return False

    def list_files(self, folder_id: Optional[str] = None) -> List[dict]:
        """Listet Dateien auf"""
        if not self.service:
            return []

        try:
            query_parts = []

            if folder_id:
                query_parts.append(f"'{folder_id}' in parents")

            # Nur PDF und Word
            mime_types = [
                'application/pdf',
                'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            ]
            mime_query = " or ".join([f"mimeType='{mt}'" for mt in mime_types])
            query_parts.append(f"({mime_query})")

            query = " and ".join(query_parts)

            results = self.service.files().list(
                q=query,
                pageSize=1000,
                fields="files(id, name, mimeType, createdTime, modifiedTime)"
            ).execute()

            return results.get('files', [])

        except Exception as e:
            self.logger.error(f"Fehler beim Auflisten: {e}")
            return []

    def download_file(self, file_id: str) -> Optional[bytes]:
        """Lädt Datei herunter"""
        if not self.service:
            return None

        try:
            request = self.service.files().get_media(fileId=file_id)
            file_content = io.BytesIO()
            downloader = MediaIoBaseDownload(file_content, request)

            done = False
            while not done:
                status, done = downloader.next_chunk()

            file_content.seek(0)
            return file_content.read()

        except Exception as e:
            self.logger.error(f"Download fehlgeschlagen: {e}")
            return None


class DataCollectionService:
    """
    Hauptservice für Datensammlung

    Implementiert Phase 2 des CRISP-DM: Data Understanding
    """

    def __init__(self, config):
        self.config = config
        self.logger = logging.getLogger(__name__)

        self.file_reader = FileReader()
        self.gdrive = GoogleDriveConnector()

        # Statistiken
        self.stats = {
            'total_collected': 0,
            'successful': 0,
            'failed': 0,
            'by_source': {},
            'by_format': {}
        }

    def collect_local_files(self, directory: Path) -> List[JobAd]:
        """Sammelt lokale Dateien"""
        self.logger.info(f"📂 Sammle lokale Dateien aus: {directory}")

        if not directory.exists():
            self.logger.error(f"Verzeichnis nicht gefunden: {directory}")
            return []

        job_ads = []

        # Dateien finden
        patterns = ['*.pdf', '*.docx', '*.doc']
        files = []

        for pattern in patterns:
            if self.config.recursive_search if hasattr(self.config, 'recursive_search') else True:
                files.extend(directory.rglob(pattern))
            else:
                files.extend(directory.glob(pattern))

        self.logger.info(f"   Gefunden: {len(files)} Dateien")

        # Verarbeiten
        for i, file_path in enumerate(files, 1):
            self.logger.info(f"   [{i}/{len(files)}] {file_path.name}")

            job_ad = self._process_local_file(file_path)
            if job_ad:
                job_ads.append(job_ad)
                self.stats['successful'] += 1
            else:
                self.stats['failed'] += 1

            self.stats['total_collected'] += 1

        self.stats['by_source']['local'] = len(job_ads)

        return job_ads

    def _process_local_file(self, file_path: Path) -> Optional[JobAd]:
        """Verarbeitet einzelne lokale Datei"""
        try:
            # Text extrahieren
            text = self.file_reader.read_file(file_path)

            if not text or len(text) < 100:
                self.logger.warning(f"      ⚠️  Zu wenig Text extrahiert")
                return None

            # JobAd erstellen
            job_ad = self._create_job_ad(
                file_name=file_path.name,
                raw_text=text,
                source="local"
            )

            # Format-Statistik
            fmt = file_path.suffix.lower()
            self.stats['by_format'][fmt] = self.stats['by_format'].get(fmt, 0) + 1

            self.logger.info(f"      ✓ {len(text)} Zeichen")
            return job_ad

        except Exception as e:
            self.logger.error(f"      ✗ Fehler: {e}")
            return None

    def collect_google_drive(self, folder_id: Optional[str] = None) -> List[JobAd]:
        """Sammelt Dateien von Google Drive"""
        self.logger.info("☁️  Sammle Google Drive Dateien")

        if not GDRIVE_AVAILABLE:
            self.logger.error("Google Drive nicht verfügbar")
            return []

        # Authentifizieren
        if not self.gdrive.authenticate():
            return []

        # Dateien auflisten
        files = self.gdrive.list_files(folder_id)
        self.logger.info(f"   Gefunden: {len(files)} Dateien")

        job_ads = []

        for i, file in enumerate(files, 1):
            self.logger.info(f"   [{i}/{len(files)}] {file['name']}")

            job_ad = self._process_gdrive_file(file)
            if job_ad:
                job_ads.append(job_ad)
                self.stats['successful'] += 1
            else:
                self.stats['failed'] += 1

            self.stats['total_collected'] += 1

        self.stats['by_source']['google_drive'] = len(job_ads)

        return job_ads

    def _process_gdrive_file(self, file: dict) -> Optional[JobAd]:
        """Verarbeitet einzelne Google Drive Datei"""
        try:
            # Download
            file_bytes = self.gdrive.download_file(file['id'])
            if not file_bytes:
                self.logger.warning(f"      ⚠️  Download fehlgeschlagen")
                return None

            # Text extrahieren
            file_obj = io.BytesIO(file_bytes)

            if file['name'].endswith('.pdf'):
                text = self.file_reader.read_pdf(file_obj, is_bytes=True)
            elif file['name'].endswith(('.docx', '.doc')):
                text = self.file_reader.read_word(file_obj, is_bytes=True)
            else:
                return None

            if not text or len(text) < 100:
                self.logger.warning(f"      ⚠️  Zu wenig Text")
                return None

            # JobAd erstellen
            job_ad = self._create_job_ad(
                file_name=file['name'],
                raw_text=text,
                source="google_drive"
            )

            self.logger.info(f"      ✓ {len(text)} Zeichen")
            return job_ad

        except Exception as e:
            self.logger.error(f"      ✗ Fehler: {e}")
            return None

    def _create_job_ad(self, file_name: str, raw_text: str, source: str) -> JobAd:
        """Erstellt JobAd-Objekt aus Rohdaten"""

        # Eindeutige ID generieren
        unique_id = hashlib.md5(
            (file_name + raw_text[:100]).encode()
        ).hexdigest()[:12]

        # Basis-JobAd
        job_ad = JobAd(
            id=unique_id,
            file_name=file_name,
            source=source,
            raw_text=raw_text,
            char_count=len(raw_text),
            word_count=len(raw_text.split()),
            collection_date=datetime.now()
        )

        # Basis-Extraktion (schnell)
        job_ad.job_title = self._extract_title_quick(file_name)

        return job_ad

    def _extract_title_quick(self, file_name: str) -> str:
        """Schnelle Titel-Extraktion aus Dateiname"""
        title = file_name.replace('.pdf', '').replace('.docx', '').replace('.doc', '')
        title = re.sub(r'[_-]', ' ', title)
        title = re.sub(r'\d{2}\.\d{2}\.\d{2,4}', '', title)
        return title.strip()[:100] or "Unbekannt"

    def get_statistics(self) -> dict:
        """Gibt Sammel-Statistiken zurück"""
        return self.stats
